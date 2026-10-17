
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document

OUTPUT_DIR = Path(__file__).parent / "data" / "sample_docs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# 1. PDF sample: a short company policy document.
pdf_path = OUTPUT_DIR / "company_policies.pdf"
pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
pdf.setTitle("Company Policies")

pdf.drawString(72, 740, "Company Policies")
pdf.drawString(72, 710, "Section 1: Leave Approval")
pdf.drawString(
    72, 685,
    "Employees must obtain manager approval before planned leave."
)
pdf.drawString(
    72, 660,
    "Emergency leave should be reported to the manager promptly."
)
pdf.drawString(72, 620, "Section 2: IT Support")
pdf.drawString(
    72, 595,
    "Contact the IT Support team when company equipment fails."
)
pdf.save()

# 2. HTML sample: a product page with navigation and footer noise.
html_path = OUTPUT_DIR / "product_page.html"
html_path.write_text(
    """<!doctype html>
<html>
<head><title>Enterprise Laptop</title></head>
<body>
<nav>Home | Products | Support | Contact</nav>
<main>
  <h1>Enterprise Laptop</h1>
  <h2>Product Overview</h2>
  <p>A business laptop designed for everyday office workloads.</p>
  <h2>Warranty</h2>
  <p>The product includes a standard one-year limited warranty.</p>
</main>
<footer>Copyright 2026 Example Company</footer>
</body>
</html>
""",
    encoding="utf-8",
)

# 3. DOCX sample: headings and paragraphs.
docx_path = OUTPUT_DIR / "onboarding.docx"
doc = Document()
doc.add_heading("Employee Onboarding Guide", level=1)
doc.add_heading("First Day", level=2)
doc.add_paragraph(
    "Meet your manager and complete the required orientation."
)
doc.add_heading("IT Setup", level=2)
doc.add_paragraph(
    "Request approved access to the tools required for your role."
)
doc.save(docx_path)

print("Sample documents created:")
for path in sorted(OUTPUT_DIR.iterdir()):
    if path.is_file():
        print(f"- {path.name}: {path.stat().st_size} bytes")
